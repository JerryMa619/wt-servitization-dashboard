from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from pathlib import Path
ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'docs/cmapss/C-MAPSS_English_Viva_Presentation_Script.docx'
d=Document(); sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(2);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2.1);sec.right_margin=Cm(2.1)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
 s=d.styles[name];s.font.name='Arial';s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Arial Unicode MS')
 s.font.color.rgb=RGBColor.from_string('000000')
for st in d.styles:
 for el in list(st._element.iter(qn('w:pBdr'))):el.getparent().remove(el)
 for el in st._element.iter(qn('w:rFonts')):
  for key in list(el.attrib):
   if 'Theme' in key:del el.attrib[key]
normal=d.styles['Normal'];normal.font.size=Pt(10.5);normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.22
for name,size in [('Title',25),('Heading 1',18),('Heading 2',12)]:
 s=d.styles[name];s.font.size=Pt(size);s.font.bold=True;s.paragraph_format.space_before=Pt(12);s.paragraph_format.space_after=Pt(8)
 s.paragraph_format.keep_with_next=True
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('C-MAPSS English Viva Script  |  ');r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
def p(text,style=None):return d.add_paragraph(text,style)
def h(text):return d.add_heading(text,2)
def page(title):d.add_page_break();d.add_heading(title,1)
def table(headers,rows,widths):
 t=d.add_table(rows=1, cols=len(headers));t.autofit=False
 borders=OxmlElement('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'D9D9D9');borders.append(edge)
 t._tbl.tblPr.append(borders)
 for c,w in zip(t.columns,widths):c.width=Cm(w)
 for i,v in enumerate(headers):t.rows[0].cells[i].text=v
 rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
 for row in rows:
  cells=t.add_row().cells
  for i,v in enumerate(row):cells[i].text=str(v)
 for ri,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  for i,c in enumerate(row.cells):
   c.width=Cm(widths[i]);tcpr=c._tc.get_or_add_tcPr()
   sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E4EEE9' if ri==0 else ('F5F7F4' if ri%2 else 'FFFFFF'));tcpr.append(sh)
   margins=OxmlElement('w:tcMar')
   for side in ['top','bottom','left','right']:
    x=OxmlElement('w:'+side);x.set(qn('w:w'),'40' if side in ['top','bottom'] else '90');x.set(qn('w:type'),'dxa');margins.append(x)
   tcpr.append(margins)
   for par in c.paragraphs:
    par.paragraph_format.space_after=Pt(3);par.paragraph_format.line_spacing=1.13
    for run in par.runs:run.font.size=Pt(9);run.bold=ri==0
 d.add_paragraph().paragraph_format.space_after=Pt(0)
 return t
def link(label,url):
 par=d.add_paragraph();r=OxmlElement('w:hyperlink');r.set(qn('r:id'),par.part.relate_to(url,RT.HYPERLINK,is_external=True))
 run=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'17665E');pr.append(color);run.append(pr);tx=OxmlElement('w:t');tx.text=label;run.append(tx);r.append(run);par._p.append(r)

import json

def cue(text):
 par=p('Demonstration cue: '+text)
 for r in par.runs:r.italic=True;r.font.color.rgb=RGBColor.from_string('52635C')
def note(text):
 par=p('Speaker note — not spoken: '+text)
 for r in par.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string('52635C')

p('C-MAPSS Digital Twin Servitization', 'Title')
p('English Viva Presentation Script', 'Subtitle')
p('English edition of the v0.8 companion | 7 October 2026')
h('Purpose and central argument')
p('This script explains how the demonstration connects degradation evidence, service responsibility and proposed advice. It combines a short opening speech, scene-by-scene narration and evidence for viva questions. The technical results and limitations follow the reviewed v0.8 demonstration.')
p('The research question is: how can technical condition and service responsibility be brought into one inspectable chain of decision evidence? The prototype implements observation replay, prognosis, contract assumptions, advice and semantic checks. Authorised maintenance delivery and observed service outcomes remain outside its implemented scope.')
table(['Layer','What the demonstration contains'],[
['Asset evidence','Original C-MAPSS benchmark observations, cycle windows, remaining useful life estimates and empirical residual intervals.'],
['Service context','Assumed contracts, consequence costs, intervention thresholds and proposed recommendations.'],
['Evidence handoff','RDF relationships, executable SPARQL queries, SHACL checks, model identity and data provenance.']],[3.5,13.3])
h('How to use this script')
p('Read the three-minute overview for a short introduction. Use the scene scripts for a longer walkthrough, then select the evidence sections in response to questions. Demonstration cues and speaker notes are not part of the spoken text. All presentation content is in English; three optional Chinese notes highlight distinctions that are easy to overstate.')
link('Open the English demonstration','https://jerryma619.github.io/wt-servitization-dashboard/cmapss/?story=1')
p('Select Automatic overview for the nine-scene animation or Viva · manual inspection to pause and inspect individual chapters. The automatic sequence lasts approximately three minutes, with additional time possible for semantic execution.')

page('1 Three minute opening script')
cue('Start the automatic overview. Use this speech as an overview; do not force each paragraph to follow the animation timer.')
p('My research examines how digital twin functions and ontology-based data structures can represent service responsibility around high-value assets. In this case study, I ask how degradation evidence can support a service recommendation whose assumptions and provenance can be inspected.')
p('The demonstration replays the C-MAPSS simulation benchmark. At each displayed cycle, the prediction uses only the observations available up to that point. A remaining useful life estimate describes the technical evidence. The service recommendation also depends on assumed contract responsibilities, costs and intervention rules.')
p('The digital twin framework locates the responsibilities for collecting observations, producing estimates and forming advice. The ontology connects the engine, observation window, estimate, model, contract and recommendation. Together, they make it possible to trace a recommendation back to its evidence. The browser executes semantic queries and constraint checks on those relationships.')
p('There are two useful comparisons in the story. At cycle 158, changing the assumed consequence cost changes the action ranking while neither intervention threshold is active. At cycle 171, a stricter threshold changes which actions are eligible. These comparisons show why the same technical evidence can lead to different advice under different service assumptions.')
p('The evaluation also exposes limitations. Widening the uncertainty interval changes some recommendations, and the frequency depends on which observation window is examined. The intervals are capped at 125 cycles, so they cannot cover true remaining lives above that cap. In a bounded comparison, conventional JSON checks and semantic checks produce equivalent results. I therefore claim executable evidence checking, rather than demonstrated ontology superiority.')
p('The micro wind turbine case complements this demonstration by showing how a simulated service process operates, including downtime, repair and state updates. C-MAPSS concentrates on how a recommendation is justified. Together, the cases support the feasibility and inspectability of the method. Real service benefits, wider interoperability and the precise research contribution still require independent evaluation and comparison with prior work.')
note('中文提示：proposed recommendation 是拟议建议，不等于已批准或已执行的维护。')

page('2 The roles of the two case studies')
p('I use the two cases to examine different parts of digital twin servitization. Both use architectural responsibilities and semantic relationships. Their distinction lies in the process and evidence that each case can demonstrate.')
table(['Dimension','Micro wind turbine','C-MAPSS'],[
['Main question','How does the service process operate?','How is a recommendation justified?'],
['Narrative','Simulated condition, service selection, downtime and repair, then state update.','Degradation observations, prefix-based prognosis, contract context, then advice and evidence.'],
['Framework emphasis','Coordination of monitoring, decisions, simulated delivery and feedback.','Responsibilities and handoffs between data, prognosis and advice.'],
['Ontology emphasis','Links between operating snapshots, candidate actions, model features and service evidence.','Asset identity, cycle units, prediction provenance, contract policy and recommendation links.'],
['Evaluation emphasis','Simulated service accounting, policy comparisons and parameter sensitivity.','Prediction uncertainty, fixed-evidence comparisons, traceability and recommendation sensitivity.'],
['Evidence boundary','Simulated or reference-assisted condition, pseudo-hour life and assumed economics; no independent field intervention validation.','Benchmark replay without real contracts, repair trajectories or measured service benefits.']],[3.1,6.85,6.85])
h('Suggested spoken transition')
p('The wind turbine case shows how I organise a simulated service process. I then use C-MAPSS to inspect the basis of a recommendation: which observations support it, how service responsibility enters the policy, and how uncertainty affects the result. This is a change in evaluation emphasis within the same research approach.')
p('Across the cases, I reuse responsibility groups and evidence relationships, while adapting units, estimate types and action assumptions. That reuse demonstrates implementation across two settings. It does not establish universal applicability, and simulated repair benefits from the wind turbine cannot be reported as realised benefits in C-MAPSS.')

page('3 Scene scripts for observations and prognosis')
h('Scene 1 The research object')
cue('Point to asset evidence, service context and service delivery.')
p('I begin by separating the asset evidence from the service context. The engine observations describe degradation. The service context specifies assumed responsibilities and rules. The prototype connects those two layers to proposed advice. Physical service delivery and its outcomes are not implemented in this case.')
h('Scene 2 Recorded observations')
cue('Follow FD001 Engine 034 through cycles 30 to 94.')
p('These are recorded benchmark observations, replayed more quickly for the presentation. Each observation retains its engine identity, cycle and source. The animation represents progression through the record; it is not a live connection to a physical engine.')
h('Scene 3 Prognosis and uncertainty')
cue('Follow the estimate to cycle 158 and identify the point estimate and interval.')
p('The prognosis uses a feature window of up to 30 available cycles. The model was fitted offline, and the browser displays its exported results. At a given replay cycle, future observations are not used to form that cycle’s estimate.')
p('Remaining useful life is expressed in cycles. The interval comes from held-out training residual quantiles and is clipped to the model’s range. It is an empirical representation of uncertainty, not a guarantee. A health classification and a service recommendation are also different outputs: the recommendation applies additional service assumptions.')
h('What the audience should retain')
p('The unit of evidence is a traceable snapshot, not an isolated number. Its interpretation requires the engine, observation window, model and units. The selected story follows one engine to explain the mechanism; it does not represent every engine or operating condition.')

page('4 Scene scripts for the framework and ontology')
h('Scene 4 Responsibilities and evidence handoffs')
cue('Pause at cycle 158. Follow the framework highlights and wait for the actual query and validation results.')
p('The framework tells me where each responsibility sits. The ontology specifies the identities, units and relationships carried between those responsibilities. My mapping is inspired by ISO 23247; it is not a claim of complete conformance to the standard.')
table(['Group','Implementation and exchanged objects'],[
['OE','The observable asset boundary: engine identity, dataset and cycle context. No physical engine is connected.'],
['DCE','The replay adapter supplies Observation and ObservationWindow objects, with sensor, cycle and source information.'],
['DTE','Window features and exported model results supply a CycleRULEstimate linked to its Model.'],
['UE','The TypeScript advice function combines the estimate with Contract and Policy objects to produce a Recommendation.'],
['CS','RDF materialisation, SPARQL, SHACL and provenance export support inspection of the resulting evidence.']],[2.0,14.8])
p('The actual sequence matters. I first obtain observations and prognosis, then compute the recommendation in TypeScript. I next materialise the evidence as RDF and run the query and constraint checks. The ontology does not infer the action through OWL, and the checks do not implement an approval gate.')
h('Scene 5 A missing unit in a test copy')
cue('Run the missing-unit example and identify the nonconforming SHACL report.')
p('Here I deliberately remove the cycle unit from a test copy of the estimate. SHACL reports that the evidence fails the required constraint. This demonstrates an executable check on the representation. It does not indicate a new engine fault, and a well-designed conventional validator can detect the same defect.')
p('The collaboration between the framework and ontology is therefore concrete: responsibilities have implementation locations, and their exchanged objects can be traced through explicit relationships. The animated highlights explain that sequence; they are not evidence of live network communication.')
note('中文提示：SHACL 检查表达与关系约束；不证明预测准确、行动安全或合同履约。')

page('5 Scene scripts for service responsibility and advice')
h('Scene 6 Consequence cost changes the ranking')
cue('Compare the two service contexts at cycle 158.')
p('I hold the technical evidence fixed. With an assumed consequence cost of five, the model recommends Continue; with twenty, it recommends Enhanced Monitoring. Neither intervention threshold is active here. This comparison isolates how a change in assumed exposure can change the ranking of actions.')
h('Scene 7 The threshold changes eligibility')
cue('Move to cycle 171 and open the cost-by-threshold comparison.')
p('The point estimate is 45.0418 cycles and the lower bound is 20.139 cycles. The lower bound lies between the thresholds of 15 and 25 cycles. The four combinations separate the effect of cost from the effect of the intervention rule.')
table(['Consequence cost','Threshold in cycles','Computed recommendation'],[['5','15','Enhanced Monitoring'],['20','15','Inspection'],['5','25','Planned Maintenance'],['20','25','Planned Maintenance']],[4.2,4.4,8.2])
p('At a threshold of 15, increasing cost changes Monitoring to Inspection. At either cost, raising the threshold to 25 leads to Planned Maintenance. Maintenance cost and the waiting-period multiplier remain at one. These are analytical scenarios, not four observed commercial contracts. They explain the implemented mechanism, not a causal effect on service outcomes.')
h('Scene 8 Proposed evidence handoff')
cue('Move to cycle 181 and show the evidence export.')
p('By this point the reference policy also proposes maintenance. The export preserves the recommendation and its supporting estimate, policy, model and source relationships. A planner would still need to review the evidence and authorise an action. No actual approval, scheduling or maintenance outcome is recorded here.')
h('Scene 9 Contributions and limits')
p('The implemented result is an inspectable path from observations to advice. The case supports responsibility mapping, semantic evidence checking and analysis of contract assumptions. I distinguish those demonstrated capabilities from claims about industrial benefit, generalisation or theoretical novelty.')

page('6 The decision model and service accountability')
h('Action scores and their assumptions')
p('I score each candidate action as its direct cost plus an assumed consequence cost multiplied by a normal-approximation probability over the action’s waiting period. Continue, Enhanced Monitoring and Inspection use illustrative waiting periods of 20, 10 and 5 cycles. Planned Maintenance and Derate/Hold use zero.')
p('Zero waiting time removes the modelled pre-intervention exposure term. It does not remove execution risk or residual risk. Enhanced Monitoring has no explicit information-gain model. The score is therefore a transparent mechanism for exploring assumptions, rather than an operationally validated optimum.')
h('Contract targets and decision rules')
p('The availability-assurance scenario uses consequence cost 20 and an intervention threshold of 25 cycles. The maintenance-support scenario uses cost 5 and threshold 15. These assumptions represent different responsibility exposures; they are not derived from the headline availability targets.')
p('The illustrative targets are 99% and 95% across 200 planned service opportunities. They allow losses of two and ten opportunities respectively. If planned maintenance is assumed to consume five opportunities, the remaining budgets are −3 and +5. This separate budget calculation does not enter the recommendation optimiser. It is neither measured time availability nor evidence that a contract has been fulfilled.')
table(['Responsibility stage','Present interpretation'],[
['Operator requirement','Service continuity is the illustrative requirement; the target and opportunity count are assumed.'],
['Provider exposure','Consequence cost and threshold encode explicit research assumptions.'],
['Policy application','The advice function applies those assumptions to a fixed estimate and retains the evidence links.'],
['Planner review','A future reviewer could accept, reject or request more evidence. Authenticated authority and approval are not implemented.'],
['Service delivery','Scheduling, intervention and outcome observation remain unimplemented; the recommendation is proposed.']],[4.2,12.6])

page('7 Prediction evidence and the interval cap')
p('The detailed replay covers 32 engines and 4,835 snapshots. Prediction metrics use the complete test endpoints of each subset, which is a different evaluation population. Each subset has its own fitted model; this is not a cross-subset transfer experiment.')
current=json.loads((ROOT/'src/cmapss/evaluation-evidence.json').read_text())
table(['Subset','Endpoint RMSE in cycles','Overall empirical coverage'],[['FD001','16.69','78 / 100  (78.00%)'],['FD002','28.94','174 / 259  (67.18%)'],['FD003','17.86','69 / 100  (69.00%)'],['FD004','30.87','152 / 248  (61.29%)']],[3,6.2,7.6])
p('Training targets are capped at 125 cycles. Evaluation retains uncapped endpoint truth. The residual 10th and 90th quantiles define an 80% reference interval, but empirical test coverage is lower and is not guaranteed. Using a normal approximation for action scoring introduces a further assumption.')
h('Why the cap matters')
table(['Subset','All endpoints covered','Truth at most 125','Truth above 125'],[[r['dataset'],f"{r['all']['covered']} / {r['all']['n']}",f"{r['withinCap']['covered']} / {r['withinCap']['n']}",f"{r['aboveCap']['covered']} / {r['aboveCap']['n']}"] for r in current['intervals']],[3,4.6,4.6,4.6])
p('Because the upper bound is clipped to 125, no endpoint with true remaining life above 125 can be covered. Coverage within the cap is approximately 87.64%, 86.14%, 81.18% and 83.98% respectively. These conditional results explain part of the limitation; they do not replace the overall results or establish calibration. Mean interval widths are 39.2, 39.8, 36.0 and 41.9 cycles.')
p('The FD001 endpoints were reconstructed using the frozen fit and training residual quantiles, checked against source hashes, the eight published replay endpoints and aggregate metrics. Test labels were not used to tune that fit. However, the test set has now been inspected diagnostically; any redesign needs a new, untouched evaluation set.')

page('8 Sensitivity and engine level evidence')
p('For this comparison I use availability assurance as the reference contract and change one parameter at a time. The denominator is 1,233 snapshots from eight selected FD001 engines. These repeated snapshots provide descriptive counts, not independent experimental success rates.')
table(['Scenario','Advice changed','Different from threshold comparator'],[['Reference','0','121'],['Consequence cost ×0.75','8','113'],['Consequence cost ×1.25','21','135'],['Threshold −5 cycles','26','147'],['Threshold +5 cycles','40','81'],['Waiting period ×0.75','43','78'],['Waiting period ×1.25','70','168'],['Interval width ×0.75','108','67'],['Interval width ×1.25','150','165']],[6.5,3.7,6.6])
p('The threshold comparator selects Planned Maintenance when the lower bound is at or below the same threshold, and Continue otherwise. Disagreement does not establish better advice. Width perturbations hold the point estimate fixed, scale its distances to the bounds, and clip to 0–125 cycles; they do not recalibrate the model.')
h('Widening the interval by a factor of 1.25')
analysis=json.loads((ROOT/'docs/cmapss/comparison-evidence/decision-analysis.json').read_text())
row=analysis['sensitivity'][1]['rows'][8]
rows=[]
for e in row['byEngine']:
 fmt=lambda x:'Not observed' if x is None else str(x)
 rows.append([str(e['engine']).zfill(3),f"{e['changed']} / {e['frames']}",fmt(e['referenceFirst'])+' → '+fmt(e['scenarioFirst']),fmt(e['firstCycleShift'])])
table(['Engine','Changed / snapshots','First maintenance advice','Cycle shift'],rows,[2,4.1,7.1,3.6])
p('The snapshot-weighted change rate is 150/1,233, or 12.17%; the equal-engine mean is 10.68%. Changed actions comprise 92 Continue-to-Monitoring, 48 Monitoring-to-Maintenance and 10 Monitoring-to-Inspection transitions. Unchanged counts are 852 Continue, 63 Monitoring and 168 Maintenance.')
p('First advice is measured only within the supplied replay. Negative shifts indicate earlier advice, not executed maintenance. Engine 001 contributes only two snapshots; equal weighting cannot remove unequal observation windows or establish representativeness.')

page('9 Matched windows and mechanism diagnostics')
h('A common observation window')
p('I also examine cycles 30–120 inclusive for the seven engines with a record at every cycle: 020, 034, 049, 068, 081, 090 and 100. Each contributes 91 snapshots, giving 637 in total. Engine 001 is excluded because it provides only two of the required cycles.')
table(['Availability assurance','Full replay','Matched cycles 30–120'],[['Sample','8 engines; 1,233 snapshots','7 engines; 637 snapshots'],['Advice changes at width ×1.25','150 / 1,233  (12.17%)','4 / 637  (0.63%)']],[5.8,5.5,5.5])
p('Maintenance support has zero changes in the matched window. Equal-engine and snapshot weighting coincide because each included engine supplies 91 observations. This is a retrospective complete-case diagnostic, not a preregistered or representative sample. Equal cycles do not imply equal degradation stages.')
p('The earlier window omits later degradation. Its lower change rate demonstrates dependence on observation scope, rather than improved full-life robustness. I therefore retain the full replay result alongside it.')
h('Separating the two uses of interval width')
table(['Score width','Eligibility width','Maintenance support changes','Availability assurance changes'],[['1×','1×','0','0'],['1×','1.25×','31','48'],['1.25×','1×','91','150'],['1.25×','1.25×','97','150']],[3,3,5.4,5.4])
p('All four rows use the same 1,233 snapshots. I vary the interval used for probability scoring separately from the interval used for threshold eligibility. The mixed rows are synthetic implementation diagnostics, not calibrated intervals suitable for deployment.')
p('The combined change count is not the sum of the separate counts. These comparisons help locate how the code produces a different recommendation; they do not estimate a causal effect on service outcomes.')
note('中文提示：较低变化率不等于效果更好；必须同时交代样本、窗口和分母。')

page('10 What the semantic comparison demonstrates')
p('I compare the same FD001 Engine 034 evidence at cycle 171 through two implementations. Conventional JSON rules and independently expressed SHACL Core constraints check the same scalar requirements. A JavaScript join and a SPARQL query answer the same provenance question. This bounded test profile is separate from the full production ontology.')
table(['Rule','Shared requirement'],[['R1','The estimate explicitly uses cycles.'],['R2','The estimate’s asset matches the snapshot’s asset.'],['R3','The recommendation’s contract version matches the supplied expected version.'],['R4','The estimate retains a source hash of 64 lower-case hexadecimal characters.'],['R5','The recommendation links to the estimate for this snapshot.']],[2,14.8])
p('The seven demonstration fixtures are valid evidence, a missing unit, an hours unit, an incorrect asset, an outdated contract version, missing provenance and a missing estimate link. Both implementations return the expected checks, and their query answers agree. Invalid evidence can still be queryable; a query result is not an approval.')
h('Additional regression scenarios')
p('With the validators unchanged, I add 35 project-authored scenarios: all 32 combinations of five defect categories, plus a consistent contract-version update, a stale recommendation after a version update, and a consistent identity migration. Both implementations satisfy the expected checks and return equivalent query results. Additional checks also cover a non-empty malformed hash.')
p('This result supports functional parity on the specified tasks. It does not show ontology superiority, faster execution, independently validated interoperability or lower change effort. The scenarios were authored within the project, and the expected contract version is supplied as input rather than retrieved from an authoritative live system.')
h('Why retain an equal result')
p('The equal result places the research claim at the appropriate level. I have implemented explicit semantic relationships, queries and constraints that can be inspected and exported. To establish an advantage over a well-designed conventional approach, I need an independent comparison on tasks where the proposed advantage can be measured.')

page('11 Contributions and the remaining evaluation')
p('My candidate contribution is a method for mapping service responsibility to digital twin functions and organising recommendation evidence through explicit identities, units and provenance. The cases and comparisons make that implementation open to examination. The contribution still needs to be distinguished from prior work.')
table(['Candidate contribution','Present evidence','Further evaluation needed'],[
['Responsibility mapping','Located functions and handoff objects connect observations, prognosis, policy and advice.','Architectural comparison, suitability of the allocation and independently assessed reuse.'],
['Inspectable semantic evidence','Executed queries and constraints, provenance exports and bounded JSON–RDF parity.','Independent competency questions, external schemas, unseen defects and measured change effort.'],
['Advice in a contract context','Fixed-evidence factorial comparisons, engine-level sensitivity and mechanism diagnostics.','Empirical parameters, untouched evaluation data and observed service outcomes.']],[4,6.2,6.6])
h('Independent semantic evaluation protocol')
p('Before adapting either implementation, an independent reviewer should specify the business questions and expected answers. The reviewer should then provide withheld defects and an external source structure. Both approaches should receive the same inputs, requirements and change tasks. The comparison should retain correctness, lost mappings, edits and reviewer effort. This protocol has been specified but not carried out.')
h('Data required for service outcome evaluation')
p('A future study needs an anonymous asset identifier; contract version and authority; the basis of responsibility and costs; decision time and then-available evidence; proposed action; authorised action and reviewer; actual intervention and time; measured downtime or exposure; observed outcomes; and a comparable baseline with relevant confounders.')
p('The population, observation period, service outcome and comparator should be defined before outcomes are collected. Proposal, authorisation, execution and outcome are separate events. A lower score under the same assumed cost function cannot independently demonstrate real operational benefit.')
h('Current review judgement')
p('The matched-window analysis, interval-cap diagnosis and semantic regression tests strengthen the inspectability of the demonstration. The remaining priorities are external evaluation, fresh calibration and validation splits, empirical service evidence and a full-text comparison with prior research. Engineering acceptance tests establish that the demonstration works as designed; they do not validate its scientific hypotheses.')

page('12 Viva questions and suggested answers')
h('Is this a complete digital twin')
p('I describe it as a replay-driven evidence-to-advice prototype. It implements a defined set of digital twin responsibilities around benchmark observations. It has no live bidirectional physical connection or completed physical service loop.')
h('Why use an ontology rather than a database')
p('The ontology provides explicit relationships that I can query, validate and export. My bounded comparison shows parity with conventional checks. A wider advantage would require independent tasks involving heterogeneous sources, change and reuse; I do not claim that advantage from the current fixtures.')
h('Does semantic validation improve prediction accuracy')
p('It does not directly improve accuracy. It checks the representation and its relationships. Prediction calibration, action suitability and contract fulfilment each require separate evidence. A conforming record may still contain a poor estimate.')
h('Do the contract assumptions determine the answer')
p('They influence the answer, which is why I state them explicitly and vary them while holding the technical evidence fixed. The factorial and sensitivity analyses expose their effects. They do not establish that the assumed values reflect a real contract or produce better outcomes.')
h('What is the practical significance')
p('The prototype makes a recommendation’s basis visible: the data window, model, uncertainty, contract assumptions and policy can be inspected together. That is an implemented capability relevant to review and accountability. Reduced downtime, lower real cost or better service performance remain outcomes to be tested.')
h('What is original and what remains to be established')
p('Digital twin frameworks, service-oriented twins and ontology validation already exist. My candidate contribution concerns the specific mapping of service responsibility to implemented functions and inspectable advice evidence. I need a fuller comparison to establish its incremental contribution. Two cases alone do not prove novelty or universal validity.')
h('What would you do next')
p('I would prioritise independent questions and external data, evaluate redesigned intervals on untouched data, and collect authorised interventions and service outcomes against a defined comparator. Those steps address the remaining empirical limits more directly than adding further animation.')

page('13 Prior work and traceable sources')
p('The source positioning below carries forward the v0.8 review. Its reading scope is limited; it is not a systematic literature review or a novelty verdict. Absence of a feature from an abstract does not establish its absence from the wider literature.')
table(['Source and reading scope','Implication for the claim'],[
['ISO 23247-2 2021 official scope','Entity and functional reference architecture already exist. Adopting responsibilities is not itself architectural novelty; full conformance is not claimed.'],
['Longo and colleagues 2019 author abstract, deposited 2022','Service-oriented digital twins, ontology knowledge structures and two manufacturing testbeds are prior work. Their combination alone cannot establish novelty.'],
['Nguyen and colleagues 2025 abstract and introduction','An ontology framework for turbine-blade maintenance is prior work. The abstract places implementation and validation in future work; this scoped reading cannot establish general superiority of the present prototype.'],
['W3C SHACL Recommendation 2017','RDF constraint validation is an existing standard. The candidate contribution concerns service evidence requirements and their application.']],[5.2,11.6])
for label,url in [('ISO 23247-2 official scope','https://www.iso.org/standard/78743.html'),('Longo and colleagues author abstract','https://arxiv.org/abs/2206.03268'),('Nguyen and colleagues maintenance framework','https://doi.org/10.1007/978-3-031-93891-7_56'),('W3C SHACL Recommendation','https://www.w3.org/TR/shacl/')]:link(label,url)
h('Reproducible project evidence')
p('This English edition is based on the reviewed demonstration at commit fb3102fe0e30140589343922cb59198b244c08f7. The links below are pinned to that revision so later website changes do not silently alter the evidence underlying this script.')
base='https://github.com/JerryMa619/wt-servitization-dashboard/blob/fb3102fe0e30140589343922cb59198b244c08f7/'
for label,path in [('Evaluation audit and current review','docs/cmapss/EVALUATION_AUDIT.md'),('Factorial and sensitivity results','docs/cmapss/comparison-evidence/decision-analysis.json'),('Executed JSON and RDF comparison','docs/cmapss/comparison-evidence/bounded-comparison.json'),('Endpoint and matched-window evidence','src/cmapss/evaluation-evidence.json'),('Dataset scope and model evaluation','docs/cmapss/MULTI_DATASET.md'),('Nine-scene narrative','src/cmapss/story.ts'),('Wind turbine implementation and evaluation','docs/research/WT_VIVA_IMPLEMENTATION.md')]:link(label,base+path)

d.core_properties.title='C-MAPSS Digital Twin Servitization English Viva Presentation Script'
d.core_properties.subject='English narration and evidence for the reviewed v0.8 demonstration'
d.core_properties.author='Digital Twinning Servitization Research'
d.core_properties.keywords='C-MAPSS, digital twin, servitization, ontology, viva, English'
OUT.parent.mkdir(parents=True,exist_ok=True)
d.save(OUT)
print(OUT)
